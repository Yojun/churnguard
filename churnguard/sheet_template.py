"""Excel 填寫範本：標題列說明、下拉選單與格式檢查，填完直接上傳即可。

兩張資料表（填其中一張就好）：
  會員摘要 — 每位顧客一列（最簡單，多數會員系統都能匯出）
  訂單明細 — 每筆訂單一列（最準，可用機器學習預測）
「說明」工作表不會被分析；資料表只有標題列、沒有範例，避免範例被當成真實資料。
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROWS = 5000
DATE_MIN, DATE_MAX = 36526, 73050          # 2000-01-01 ～ 2099-12-31 的 Excel 序號
CHANNELS = '"line,email,app,sms"'

# (標題, 必填, 提示, 寬度, 類型)
SUMMARY_COLS = [
    ("顧客編號", True, "會員編號、Email 或手機都可以，只要能代表同一位顧客（不能重複）。", 16, "text"),
    ("最近消費日", True, "最近一次消費的日期，例如 2026-09-30。", 14, "date"),
    ("消費次數", False, "累計消費幾次（整數）。填了才能算出這位顧客平常多久買一次。", 11, "int"),
    ("累計消費金額", False, "歷史累計消費總額。用來排序哪些顧客最值得優先挽回。", 14, "money"),
    ("首次消費日", False, "首次消費或加入會員的日期。有填，才能算出每位顧客自己的購買間隔。", 14, "date"),
    ("姓名", False, "訊息個人化用，不填也能分析（會用「親愛的顧客」）。", 12, "text"),
    ("偏好品類", False, "這位顧客最常買的品類，會寫進推播訊息。", 12, "text"),
    ("聯絡管道", False, "從下拉選單選：line / email / app / sms。", 12, "channel"),
]
ORDER_COLS = [
    ("顧客編號", True, "會員編號、Email 或手機都可以，同一位顧客請用同一個編號。", 16, "text"),
    ("訂單日期", True, "下單日期，例如 2026-09-30。", 14, "date"),
    ("訂單金額", False, "這筆訂單的金額。", 12, "money"),
    ("商品類別", False, "這筆訂單的商品類別。", 12, "text"),
    ("姓名", False, "訊息個人化用，不填也能分析。", 12, "text"),
    ("聯絡管道", False, "從下拉選單選：line / email / app / sms。", 12, "channel"),
]

HEAD_REQ = PatternFill("solid", fgColor="2F6FED")
HEAD_OPT = PatternFill("solid", fgColor="667085")
WHITE_BOLD = Font(bold=True, color="FFFFFF")


def _data_sheet(wb, title, cols):
    ws = wb.create_sheet(title)
    for j, (name, required, tip, width, kind) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=j, value=name + (" *" if required else ""))
        c.fill, c.font = (HEAD_REQ if required else HEAD_OPT), WHITE_BOLD
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.comment = Comment(("【必填】" if required else "【選填】") + tip, "ChurnGuard")
        letter = c.column_letter
        ws.column_dimensions[letter].width = width
        rng = f"{letter}2:{letter}{ROWS}"
        dv = None
        if kind == "date":
            dv = DataValidation(type="date", operator="between", formula1=str(DATE_MIN), formula2=str(DATE_MAX),
                                allow_blank=True, errorTitle="日期格式不正確", error="請輸入日期，例如 2026-09-30。")
        elif kind == "int":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                                errorTitle="次數不正確", error="請輸入 1 以上的整數。")
        elif kind == "money":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                errorTitle="金額不正確", error="請輸入 0 以上的數字。")
        elif kind == "channel":
            dv = DataValidation(type="list", formula1=CHANNELS, allow_blank=True,
                                errorTitle="請從清單選擇", error="請選 line、email、app 或 sms。")
        if dv:
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add(rng)
        fmt = {"date": "yyyy-mm-dd", "money": "#,##0", "int": "0"}.get(kind)
        if fmt:
            for r in range(2, ROWS + 1):
                ws.cell(row=r, column=j).number_format = fmt
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    return ws


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "說明"
    lines = [
        ("ChurnGuard 資料填寫範本", True),
        ("", False),
        ("1. 選一張工作表來填就好（不用兩張都填）：", True),
        ("   「會員摘要」— 每位顧客一列。最簡單，多數會員系統／POS 都能匯出。只需要顧客編號與最近消費日，其他越完整越準。", False),
        ("   「訂單明細」— 每筆訂單一列。最準，可使用機器學習預測，並看出購買節奏的變化。", False),
        ("2. 藍色標題（有 *）是必填，灰色標題是選填。滑鼠移到標題上會看到說明。", False),
        ("3. 日期請用 2026-09-30 這種格式；聯絡管道請從下拉選單選。", False),
        ("4. 可以直接把你系統匯出的資料複製貼上，欄位順序可以不同，系統會自動辨識；沒有的欄位留空即可。", False),
        ("5. 填完存成 .xlsx，回到網站按「上傳檔案」。這張「說明」工作表不會被分析。", False),
        ("", False),
        ("填寫範例 — 會員摘要", True),
        ("   顧客編號 | 最近消費日 | 消費次數 | 累計消費金額 | 首次消費日 | 姓名 | 偏好品類 | 聯絡管道", False),
        ("   A001 | 2026-09-30 | 12 | 15800 | 2025-11-02 | 王小明 | 咖啡豆 | line", False),
        ("   A002 | 2026-05-18 | 3 | 2400 | 2026-01-20 | 林美玲 | 甜點 | email", False),
        ("", False),
        ("填寫範例 — 訂單明細", True),
        ("   顧客編號 | 訂單日期 | 訂單金額 | 商品類別 | 姓名 | 聯絡管道", False),
        ("   A001 | 2026-09-30 | 1280 | 咖啡豆 | 王小明 | line", False),
        ("   A001 | 2026-09-12 | 960 | 甜點 | 王小明 | line", False),
        ("", False),
        ("隱私小提醒：姓名、電話不是必要的。只留「顧客編號＋日期＋金額」就能分析；上傳前請確認你有權處理這份資料。", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=14 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    _data_sheet(wb, "會員摘要", SUMMARY_COLS)
    _data_sheet(wb, "訂單明細", ORDER_COLS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
